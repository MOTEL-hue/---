import openpyxl,json,sys,warnings,re
warnings.filterwarnings('ignore')
wb=openpyxl.load_workbook(sys.argv[1],data_only=False)
num=lambda v:isinstance(v,(int,float)) and not isinstance(v,bool)
S=lambda v:str(v).strip() if v is not None else ''
out={}
# ---- עסק: ציור בחול
# השנה של כל גוש נקבעת לפי שורת הסיכום שסוגרת אותו ("סה"כ תשפ"ו" בעמודה F), לא לפי כותרת בתחילתו:
# רק הגושים הראשונים מתחילים בכותרת שנה.
ws=wb['ציור בחול']
orders=[];exps=[];blk_o=[];blk_e=[]
for r in range(3,min(ws.max_row,400)+1):
    a,b,c,d,e,f,g,h,i,j,k=[ws.cell(r,x).value for x in range(1,12)]
    if S(a)=='סה"כ' or S(f).startswith('סה"כ'):
        year=S(f).replace('סה"כ','').strip().split('\\')[0].strip() or 'לא ידוע'
        for o in blk_o: o['year']=year
        for o in blk_e: o['year']=o.pop('_y',None) or year
        orders+=blk_o; exps+=blk_e; blk_o=[]; blk_e=[]
        continue
    if S(a)=='אירוע': continue
    if S(a) and num(b):
        gs=S(g); paid='V' in gs; part=''.join(ch for ch in gs if ch.isdigit()) if not paid and gs else ''
        note=S(e)+(' · שולם חלקית '+part if part else '')
        blk_o.append(dict(client=S(a),amount=b,cost=c if num(c) else 0,when=S(d),pay=S(f),note=note.strip(' ·'),paid=paid))
    name=S(j) or S(i)  # בשורה אחת ("רואה חשבון") השם נמצא בעמודת התאריך
    if name and num(k) and '\\' not in name and name!='סה"כ':
        yr=S(i) if re.match(r'^תש[א-ת]"[א-ת]$',S(i)) else ''  # בגושים הראשונים לכל הוצאה יש שנה משלה בעמודה I
        blk_e.append(dict(item=name,amount=k,when='' if not S(j) or yr else S(i),**({'_y':yr} if yr else {})))
for o in blk_o+blk_e: o['year']=o.pop('_y',None) or 'לא ידוע'
orders+=blk_o; exps+=blk_e
out['orders']=orders;out['bizExpenses']=exps
# ---- חובות
ws=wb['חובות']
debts=[]
for r in range(7,23):
    n=ws.cell(r,2).value
    if n and S(n): debts.append(dict(name=S(n),kind='הלוואה',balance=ws.cell(r,3).value or 0,monthly=ws.cell(r,4).value or 0,left=ws.cell(r,5).value or 0,urgency=ws.cell(r,6).value or 0))
    n=ws.cell(r,9).value
    if n and S(n): debts.append(dict(name=S(n),kind=S(ws.cell(r,8).value),balance=ws.cell(r,10).value or 0,monthly=0,left=ws.cell(r,11).value or 0,urgency=ws.cell(r,12).value or 0))
out['debts']=debts
# ---- חודשי
ws=wb['תמונת מצב חודשית']
inc=[];fixed=[];var=[];cons=[];charity=[]
for r in range(7,34):
    d,e=ws.cell(r,4).value,ws.cell(r,5).value
    if d and S(d)!='.' : inc.append(dict(name=S(d),amount=e if num(e) else 0,day=ws.cell(r,3).value if num(ws.cell(r,3).value) else None))
    h,i=ws.cell(r,8).value,ws.cell(r,9).value
    if h and r<=29+1 and S(h): fixed.append(dict(name=S(h),amount=i if num(i) else 0))
    l,m=ws.cell(r,12).value,ws.cell(r,13).value
    if l and S(l):
        (charity if r>=25 else var).append(dict(name=S(l),amount=m if num(m) else 0))
    p,q=ws.cell(r,16).value,ws.cell(r,17).value
    if p and S(p) and r<=23: cons.append(dict(name=S(p),amount=q if num(q) else 0))
out['monthly']=dict(income=inc,fixed=fixed,variable=var,consumption=cons,charity=charity)
# ---- שנתי
ws=wb['תמונת מצב שנתית']
yinc=[];yexp=[];cur={'F':None,'J':None}
for r in range(7,75):
    c,dv=ws.cell(r,3).value,ws.cell(r,4).value
    if c and S(c) and not S(c).startswith(('חייבות','סה"כ','צדקה','חודש','חנוכה','ירידה','חול')) and r<=32: yinc.append(dict(name=S(c),amount=dv if num(dv) else 0))
    for mc,nc,ac in (('F','G','H'),('J','K','L')):
        mv=ws[f'{mc}{r}'].value
        if mv and S(mv): cur[mc]=S(mv).replace('כסלוי','כסלו')
        n,a=ws[f'{nc}{r}'].value,ws[f'{ac}{r}'].value
        if n and S(n) and cur[mc]: yexp.append(dict(month=cur[mc],name=S(n),amount=a if num(a) else 0))
    n,a=ws.cell(r,15).value,ws.cell(r,16).value
    if n and S(n) and r<=49: yexp.append(dict(month='כללי',name=S(n),amount=a if num(a) else 0))
out['yearIncome']=yinc;out['yearExpense']=yexp
json.dump(out,open(sys.argv[2],'w',encoding='utf-8'),ensure_ascii=False,indent=1)
print({k:(len(v) if isinstance(v,list) else {kk:len(vv) for kk,vv in v.items()}) for k,v in out.items()})
import collections;print(collections.Counter(o['year'] for o in orders));print(collections.Counter(o['year'] for o in exps))
